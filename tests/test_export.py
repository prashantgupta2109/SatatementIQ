"""Unit tests for unified data export (CSV and Excel formats)."""

from datetime import date
import io
import openpyxl
import pandas as pd
import pytest

from app.models import Transaction, BankAccount, ClassificationMethod, ValidationStatus
from app.export.csv_exporter import (
    export_to_csv,
    export_to_csv_bytes,
    transactions_to_export_dataframe,
)
from app.export.excel_exporter import (
    export_to_excel,
    export_to_excel_bytes,
)


@pytest.fixture
def sample_transactions():
    return [
        Transaction(
            date=date(2026, 9, 1),
            description="SWIGGY FOOD ORDER",
            debit=450.50,
            credit=None,
            balance=24549.50,
            category="Food & Dining",
            classification_method=ClassificationMethod.RULE_BASED,
            confidence=1.00,
            validation_status=ValidationStatus.VALID,
        ),
        Transaction(
            date=date(2026, 9, 2),
            description="AMAZON INDIA SHOPPING",
            debit=1250.00,
            credit=None,
            balance=23299.50,
            category="Shopping",
            classification_method=ClassificationMethod.ML_MODEL,
            confidence=0.88,
            validation_status=ValidationStatus.VALID,
        ),
        Transaction(
            date=date(2026, 9, 3),
            description="SALARY CREDIT CORP",
            debit=None,
            credit=80000.00,
            balance=103299.50,
            category="Salary",
            classification_method=ClassificationMethod.RULE_BASED,
            confidence=0.95,
            validation_status=ValidationStatus.VALID,
        ),
    ]


@pytest.fixture
def sample_account():
    return BankAccount(
        bank_name="HDFC Bank",
        account_holder="John Doe",
        account_number="50100234567890",
        ifsc="HDFC0001234",
        statement_period="01/09/2026 to 30/09/2026",
    )


class TestCSVExport:
    """Tests for CSV export functionality."""

    def test_csv_dataframe_structure(self, sample_transactions):
        df = transactions_to_export_dataframe(sample_transactions)
        expected_cols = [
            "Date",
            "Description",
            "Debit",
            "Credit",
            "Balance",
            "Category",
            "Classification Method",
            "Confidence",
            "Validation Status",
        ]
        assert list(df.columns) == expected_cols
        assert len(df) == 3

    def test_csv_export_file_generation(self, tmp_path, sample_transactions):
        csv_file = tmp_path / "test_export.csv"
        out_path = export_to_csv(sample_transactions, output_path=str(csv_file))

        assert csv_file.exists()
        assert out_path == str(csv_file)

        df = pd.read_csv(csv_file)
        assert len(df) == 3
        assert df.iloc[0]["Description"] == "SWIGGY FOOD ORDER"
        assert df.iloc[0]["Category"] == "Food & Dining"
        assert df.iloc[0]["Debit"] == 450.50
        assert df.iloc[1]["Classification Method"] == "ML"

    def test_csv_export_bytes(self, sample_transactions):
        data_bytes = export_to_csv_bytes(sample_transactions)
        assert isinstance(data_bytes, bytes)
        content = data_bytes.decode("utf-8")
        assert "SWIGGY FOOD ORDER" in content
        assert "AMAZON INDIA SHOPPING" in content
        assert "SALARY CREDIT CORP" in content


class TestExcelExport:
    """Tests for Excel (.xlsx) multi-sheet export functionality."""

    def test_excel_export_file_generation(self, tmp_path, sample_transactions, sample_account):
        excel_file = tmp_path / "test_export.xlsx"
        out_path = export_to_excel(sample_transactions, sample_account, output_path=str(excel_file))

        assert excel_file.exists()
        assert out_path == str(excel_file)

        wb = openpyxl.load_workbook(excel_file)
        assert "Transactions" in wb.sheetnames
        assert "Account Details" in wb.sheetnames
        assert "Summary" in wb.sheetnames

        # Validate Transactions sheet
        ws_txns = wb["Transactions"]
        assert ws_txns.cell(row=1, column=1).value == "Date"
        assert ws_txns.cell(row=2, column=2).value == "SWIGGY FOOD ORDER"

        # Validate Account sheet
        ws_acc = wb["Account Details"]
        assert ws_acc["B2"].value == "BANK STATEMENT ACCOUNT SUMMARY"

        # Validate Summary sheet
        ws_sum = wb["Summary"]
        assert ws_sum.cell(row=1, column=1).value == "Category"

    def test_excel_export_bytes(self, sample_transactions, sample_account):
        excel_bytes = export_to_excel_bytes(sample_transactions, sample_account)
        assert isinstance(excel_bytes, bytes)
        # Verify ZIP header (PK\x03\x04) for valid xlsx format
        assert excel_bytes[:4] == b"PK\x03\x04"

        # Load from BytesIO
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        assert len(wb.sheetnames) == 3


class TestExportLogging:
    """Verifies that both export formats emit standard INFO logs."""

    def test_csv_and_excel_log_export_completed(self, caplog, tmp_path, sample_transactions, sample_account):
        import logging
        caplog.set_level(logging.INFO, logger="bank_processor")

        csv_file = str(tmp_path / "log_test.csv")
        export_to_csv(sample_transactions, output_path=csv_file)

        excel_file = str(tmp_path / "log_test.xlsx")
        export_to_excel(sample_transactions, sample_account, output_path=excel_file)

        log_messages = [rec.message for rec in caplog.records if rec.name == "bank_processor"]
        assert log_messages.count("Export completed") == 2
