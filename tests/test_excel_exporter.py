"""Unit tests for Excel exporter module."""

import io
from datetime import date
import openpyxl
import pytest

from app.models import Transaction, BankAccount, ClassificationMethod, ValidationStatus
from app.export.excel_exporter import export_to_excel, export_to_excel_bytes


@pytest.fixture
def sample_data():
    account = BankAccount(
        bank_name="HDFC Bank",
        account_holder="PRASHANT KUMAR",
        account_number="50100293847162",
        ifsc="HDFC0001234",
        statement_period="01/09/2026 to 30/09/2026",
    )
    transactions = [
        Transaction(
            date=date(2026, 9, 1),
            description="SWIGGY ORDER",
            debit=450.0,
            credit=None,
            balance=24550.0,
            category="Food & Dining",
            classification_method=ClassificationMethod.RULE_BASED,
            confidence=1.00,
            validation_status=ValidationStatus.VALID,
        ),
        Transaction(
            date=date(2026, 9, 2),
            description="SALARY CREDIT",
            debit=None,
            credit=75000.0,
            balance=99550.0,
            category="Salary",
            classification_method=ClassificationMethod.RULE_BASED,
            confidence=1.00,
            validation_status=ValidationStatus.VALID,
        ),
        Transaction(
            date=date(2026, 9, 3),
            description="AMAZON PURCHASE",
            debit=1500.0,
            credit=None,
            balance=98050.0,
            category="Shopping",
            classification_method=ClassificationMethod.ML_MODEL,
            confidence=0.88,
            validation_status=ValidationStatus.VALID,
        ),
    ]
    return account, transactions


class TestExcelExporter:
    """Verifies multi-sheet workbook generation, layout, and styling."""

    def test_three_sheets_present(self, sample_data):
        account, txns = sample_data
        excel_bytes = export_to_excel_bytes(txns, account)

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames

        assert "Transactions" in sheet_names
        assert "Account Details" in sheet_names
        assert "Summary" in sheet_names
        assert len(sheet_names) == 3

    def test_transactions_sheet_structure_and_panes(self, sample_data):
        account, txns = sample_data
        excel_bytes = export_to_excel_bytes(txns, account)

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb["Transactions"]

        # Check frozen panes
        assert ws.freeze_panes == "A2"
        # Check auto-filter is configured
        assert ws.auto_filter.ref is not None

        # Verify headers
        headers = [ws.cell(row=1, column=c).value for c in range(1, 11)]
        assert "Date" in headers
        assert "Description" in headers
        assert "Debit" in headers
        assert "Credit" in headers
        assert "Balance" in headers
        assert "Category" in headers
        assert "Confidence" in headers

        # Verify currency number format on debit
        debit_cell = ws.cell(row=2, column=3)  # Row 2, Debit column
        assert debit_cell.number_format == "#,##0.00"

    def test_account_details_sheet(self, sample_data):
        account, txns = sample_data
        excel_bytes = export_to_excel_bytes(txns, account)

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb["Account Details"]

        content_str = " ".join([str(c.value) for row in ws.rows for c in row if c.value])
        assert "HDFC Bank" in content_str
        assert "PRASHANT KUMAR" in content_str
        assert "50100293847162" in content_str
        assert "HDFC0001234" in content_str

    def test_summary_sheet_aggregations(self, sample_data):
        account, txns = sample_data
        excel_bytes = export_to_excel_bytes(txns, account)

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb["Summary"]

        assert ws.freeze_panes == "A2"
        headers = [ws.cell(row=1, column=c).value for c in range(1, 5)]
        assert headers == ["Category", "Transaction Count", "Total Debit", "Total Credit"]

        # Check total row at bottom
        last_row = ws.max_row
        total_label = ws.cell(row=last_row, column=1).value
        total_count = ws.cell(row=last_row, column=2).value
        total_debit = ws.cell(row=last_row, column=3).value
        total_credit = ws.cell(row=last_row, column=4).value

        assert total_label == "Total"
        assert total_count == 3
        assert total_debit == 1950.0  # 450 + 1500
        assert total_credit == 75000.0

    def test_export_to_excel_file(self, tmp_path, sample_data):
        account, txns = sample_data
        out_path = tmp_path / "test_statement.xlsx"
        result_path = export_to_excel(txns, account, output_path=str(out_path))

        assert out_path.exists()
        wb = openpyxl.load_workbook(result_path)
        assert len(wb.sheetnames) == 3
