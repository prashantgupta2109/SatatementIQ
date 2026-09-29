"""Professional Excel Exporter for Bank Statement Processing.

Generates a multi-sheet styled Excel workbook using openpyxl:
    1. Transactions: Complete tabular transaction record with auto-filters & frozen headers
    2. Account Details: Bank metadata and customer account identity
    3. Summary: Category-level aggregation (count, total debit, total credit)

Includes professional corporate styling:
    - Custom header styling & fills
    - Proper currency (#,##0.00) & date formatting
    - Dynamic column width auto-fitting
    - Filter ranges and frozen top rows
"""

import io
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import Transaction, BankAccount, ClassificationMethod
from app.export.csv_exporter import format_classification_method
from app.logging_config import get_logger

logger = get_logger("bank_processor")


# Corporate Style Palette
NAVY_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
TEAL_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E79")
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
REGULAR_FONT = Font(name="Calibri", size=11)

THIN_BORDER = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)
TOTAL_TOP_DOUBLE_BOTTOM = Border(
    top=Side(style="thin", color="000000"),
    bottom=Side(style="double", color="000000"),
)


def _auto_fit_columns(ws, min_width=12, max_width=50):
    """Dynamically adjust column widths based on content lengths."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        adjusted_width = max(max_len + 3, min_width)
        adjusted_width = min(adjusted_width, max_width)
        ws.column_dimensions[col_letter].width = adjusted_width


def build_transactions_sheet(ws, transactions: List[Transaction]):
    """Build the 'Transactions' sheet with filters, frozen panes, and formatting."""
    ws.title = "Transactions"
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"

    headers = [
        "Date",
        "Description",
        "Debit",
        "Credit",
        "Balance",
        "Category",
        "Classification Method",
        "Confidence",
        "Validation Status",
        "Review Status",
    ]

    ws.append(headers)

    # Style Header Row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = NAVY_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Write Data Rows
    for row_idx, txn in enumerate(transactions, start=2):
        date_val = txn.date.strftime("%Y-%m-%d") if txn.date else ""
        debit_val = float(txn.debit) if txn.debit is not None else None
        credit_val = float(txn.credit) if txn.credit is not None else None
        balance_val = float(txn.balance) if txn.balance is not None else None

        row_data = [
            date_val,
            txn.description or "",
            debit_val,
            credit_val,
            balance_val,
            txn.category or "Uncategorized",
            format_classification_method(txn.classification_method),
            round(txn.confidence, 2),
            txn.validation_status.value.upper(),
            txn.review_status,
        ]
        ws.append(row_data)

        # Style Cells
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER

            # Number formatting
            if col_idx == 1:  # Date
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in (3, 4, 5):  # Debit, Credit, Balance
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif col_idx == 8:  # Confidence
                cell.number_format = "0.00"
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in (7, 9, 10):  # Method, Status, Review
                cell.alignment = Alignment(horizontal="center")

    # Add Auto-Filter
    last_row = max(len(transactions) + 1, 1)
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"

    _auto_fit_columns(ws)


def build_account_sheet(ws, account: BankAccount):
    """Build the 'Account Details' sheet with metadata and styling."""
    ws.title = "Account Details"
    ws.views.sheetView[0].showGridLines = True

    # Title
    ws["B2"] = "BANK STATEMENT ACCOUNT SUMMARY"
    ws["B2"].font = TITLE_FONT

    headers = ["Field", "Extracted Value"]
    ws.cell(row=4, column=2, value=headers[0]).fill = TEAL_FILL
    ws.cell(row=4, column=2).font = HEADER_FONT
    ws.cell(row=4, column=3, value=headers[1]).fill = TEAL_FILL
    ws.cell(row=4, column=3).font = HEADER_FONT

    fields = [
        ("Bank Name", account.bank_name or "Not Specified"),
        ("Account Holder", account.account_holder or "Not Specified"),
        ("Account Number", account.account_number or "Not Specified"),
        ("IFSC Code", account.ifsc or "Not Specified"),
        ("Statement Period", account.statement_period or "Not Specified"),
    ]

    for idx, (label, val) in enumerate(fields, start=5):
        cell_lbl = ws.cell(row=idx, column=2, value=label)
        cell_lbl.font = BOLD_FONT
        cell_lbl.border = THIN_BORDER

        cell_val = ws.cell(row=idx, column=3, value=str(val))
        cell_val.font = REGULAR_FONT
        cell_val.border = THIN_BORDER

    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 38


def build_summary_sheet(ws, transactions: List[Transaction]):
    """Build the 'Summary' sheet aggregating count, debit, credit per category."""
    ws.title = "Summary"
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"

    headers = [
        "Category",
        "Transaction Count",
        "Total Debit",
        "Total Credit",
    ]
    ws.append(headers)

    # Header style
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = NAVY_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Aggregate by category
    summary: dict = {}
    for txn in transactions:
        cat = txn.category or "Other"
        if cat not in summary:
            summary[cat] = {"count": 0, "debit": 0.0, "credit": 0.0}
        summary[cat]["count"] += 1
        if txn.debit:
            summary[cat]["debit"] += txn.debit
        if txn.credit:
            summary[cat]["credit"] += txn.credit

    # Sort categories by total debit descending
    sorted_cats = sorted(summary.items(), key=lambda x: x[1]["debit"], reverse=True)

    curr_row = 2
    total_txns = 0
    total_debits = 0.0
    total_credits = 0.0

    for cat, stats in sorted_cats:
        ws.append([
            cat,
            stats["count"],
            stats["debit"],
            stats["credit"],
        ])

        total_txns += stats["count"]
        total_debits += stats["debit"]
        total_credits += stats["credit"]

        # Cell styles
        c_cat = ws.cell(row=curr_row, column=1)
        c_cnt = ws.cell(row=curr_row, column=2)
        c_deb = ws.cell(row=curr_row, column=3)
        c_crd = ws.cell(row=curr_row, column=4)

        for c in (c_cat, c_cnt, c_deb, c_crd):
            c.font = REGULAR_FONT
            c.border = THIN_BORDER

        c_cnt.alignment = Alignment(horizontal="center")
        c_deb.number_format = "#,##0.00"
        c_crd.number_format = "#,##0.00"
        c_deb.alignment = Alignment(horizontal="right")
        c_crd.alignment = Alignment(horizontal="right")

        curr_row += 1

    # Total Row
    ws.append([
        "Total",
        total_txns,
        total_debits,
        total_credits,
    ])

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=curr_row, column=col_idx)
        cell.font = BOLD_FONT
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_TOP_DOUBLE_BOTTOM
        if col_idx in (3, 4):
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")
        elif col_idx == 2:
            cell.alignment = Alignment(horizontal="center")

    # Auto-filter (excluding total row)
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}{curr_row - 1}"

    _auto_fit_columns(ws)


def export_to_excel(
    transactions: List[Transaction],
    account: Optional[BankAccount] = None,
    output_path: str = "statement_analysis.xlsx",
) -> str:
    """Generate and write a multi-sheet formatted Excel workbook to disk.

    Args:
        transactions: List of processed Transaction objects.
        account: Extracted BankAccount object.
        output_path: Destination path.

    Returns:
        The output path string.
    """
    if account is None:
        account = BankAccount()

    wb = openpyxl.Workbook()
    # Sheet 1: Transactions
    ws_txns = wb.active
    build_transactions_sheet(ws_txns, transactions)

    # Sheet 2: Account Details
    ws_acc = wb.create_sheet()
    build_account_sheet(ws_acc, account)

    # Sheet 3: Summary
    ws_sum = wb.create_sheet()
    build_summary_sheet(ws_sum, transactions)

    wb.save(output_path)
    logger.info("Export completed")
    return output_path


def export_to_excel_bytes(
    transactions: List[Transaction],
    account: Optional[BankAccount] = None,
) -> bytes:
    """Generate multi-sheet Excel workbook in-memory as bytes for direct web download.

    Args:
        transactions: List of processed Transaction objects.
        account: Extracted BankAccount object.

    Returns:
        bytes representing .xlsx file.
    """
    if account is None:
        account = BankAccount()

    wb = openpyxl.Workbook()
    ws_txns = wb.active
    build_transactions_sheet(ws_txns, transactions)

    ws_acc = wb.create_sheet()
    build_account_sheet(ws_acc, account)

    ws_sum = wb.create_sheet()
    build_summary_sheet(ws_sum, transactions)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
